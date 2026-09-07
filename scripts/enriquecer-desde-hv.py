#!/usr/bin/env python3
"""
Completa placa / marca / modelo / año / km desde la Hoja de vida (Excel).

Fuente fiable: título de cada hoja + primera fila de ficha + línea de alta
(placa / chasis). No usa aceite “Hyundai 10w30” ni fechas de mantenimiento.

Uso:
  python3 scripts/enriquecer-desde-hv.py            # dry-run
  python3 scripts/enriquecer-desde-hv.py --commit   # escribe en Supabase
"""
from __future__ import annotations

import json
import os
import re
import sys
import warnings
from collections import Counter
from pathlib import Path

import openpyxl
import requests

warnings.filterwarnings("ignore")

COMMIT = "--commit" in sys.argv
ROOT = Path(__file__).resolve().parents[1]
HOJA = ROOT / "Hoja de vida carros Autolujo Actualizadas 2025.xlsx"
if not HOJA.exists():
    HOJA = Path.home() / "Downloads" / "Hoja de vida carros Autolujo Actualizadas 2025.xlsx"

env: dict[str, str] = {}
for line in (ROOT / ".env.local").read_text().splitlines():
    line = line.strip()
    if "=" in line and not line.startswith("#"):
        k, v = line.split("=", 1)
        env[k] = v
URL = env["NEXT_PUBLIC_SUPABASE_URL"]
KEY = env["SUPABASE_SERVICE_ROLE_KEY"]
H = {
    "apikey": KEY,
    "Authorization": f"Bearer {KEY}",
    "Content-Type": "application/json",
    "Prefer": "return=representation",
}


def sb_get(path: str):
    r = requests.get(f"{URL}/rest/v1/{path}", headers=H, timeout=60)
    r.raise_for_status()
    return r.json()


def sb_patch(table: str, id_: str, patch: dict) -> bool:
    import time

    last_err = None
    for intento in range(5):
        try:
            r = requests.patch(
                f"{URL}/rest/v1/{table}?id=eq.{id_}",
                headers=H,
                json=patch,
                timeout=60,
            )
            if r.status_code >= 300:
                print("  ERR patch", id_, r.status_code, r.text[:200])
                return False
            return True
        except (requests.exceptions.SSLError, requests.exceptions.ConnectionError) as e:
            last_err = e
            time.sleep(1.5 * (intento + 1))
    print("  ERR patch", id_, last_err)
    return False


def canon_num(s: str) -> str:
    t = re.sub(r"[^A-Z0-9]", "", str(s).upper())
    m = re.match(r"^([A-Z]*)0*(\d+)$", t)
    if m:
        pref, n = m.group(1), str(int(m.group(2)))
        return pref + n if pref else n
    return t


def placa_ok(raw: str | None) -> str | None:
    if not raw:
        return None
    p = re.sub(r"[^A-Z0-9]", "", str(raw).upper())
    if not re.fullmatch(r"[A-Z]{1,3}\d{3,4}", p):
        return None
    if len(p) < 5:
        return None
    if re.match(r"^(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)\d+$", p, re.I):
        return None
    return p


def parse_ficha(texto: str) -> tuple[str | None, str | None, int | None]:
    """Solo señales de ficha (no aceite ni bitácora)."""
    # quitar menciones de aceite que ensucian marca
    u = re.sub(r"ACEITE\s+[^\n]{0,40}", " ", str(texto), flags=re.I)
    u = re.sub(r"\s+", " ", u.upper())

    marca = modelo = None
    anio = None

    if re.search(r"\bSOLUTO\b|KIA\s*SOL\b", u) or (
        re.search(r"\bKIA\b", u) and re.search(r"\bSOL", u)
    ):
        marca = "Kia"
        modelo = "Soluto"
    elif re.search(r"\bKIA\b", u):
        marca = "Kia"
        if re.search(r"\bSOLUTO\b|\bSOL\b", u):
            modelo = "Soluto"

    if re.search(
        r"GRAND\s*I[\s\-]?10|GRANDI10|GR[\-\s]?I10|GR[\-\s]?10\b|GRAND\s*I\-10",
        u,
    ):
        marca = marca or "Hyundai"
        modelo = modelo or "Grand i10"
    elif re.search(r"\bHYUNDAI\b", u) and re.search(r"I[\s\-]?10|GRAND", u):
        marca = "Hyundai"
        modelo = modelo or "Grand i10"

    # año pegado al modelo / título (no fechas sueltas de bitácora)
    near = re.search(
        r"(?:SOLUTO|GRAND\s*I[\s\-]?10|GRANDI10|GR[\-\s]?I?10|KIA|HYUNDAI|"
        r"NUEVO|AUTOMATICO|MANUAL|SEDAN)"
        r"[^\d]{0,16}(20(?:1[6-9]|2[0-8]))",
        u,
    )
    if near:
        anio = int(near.group(1))
    else:
        # título tipo ...2025 ENERO / NUEVO2024FEB
        m = re.search(r"(?:NUEVO)?(20(?:1[6-9]|2[0-8]))(?:\s*(?:ENE|FEB|MAR|ABR|MAY|JUN|JUL|AGO|SEP|OCT|NOV|DIC)|FEB|ENERO|MARZO|ABRIL|MAYO|JUNIO|AGOSTO|OCTUB)?", u)
        if m and ("NUEVO" in u or "GRAND" in u or "KIA" in u or "SOLUTO" in u or "I10" in u.replace(" ", "")):
            anio = int(m.group(1))

    return marca, modelo, anio


def extract_placa(text: str) -> str | None:
    t = str(text)
    scored: list[tuple[int, str]] = []
    for m in re.finditer(r"#\s*[A-Z0-9\-]+\s*[-–—:]\s*([A-Za-z]{1,3}\d{3,4})", t, re.I):
        p = placa_ok(m.group(1))
        if p:
            scored.append((0, p))
    for m in re.finditer(r"(?:👉🏼|👉)\s*([A-Za-z]{1,3}\d{3,4})", t):
        p = placa_ok(m.group(1))
        if p:
            scored.append((0, p))
    for m in re.finditer(r"PLACA(?:\s+PARTICULAR)?\s+([A-Za-z]{1,3}\d{3,4})", t, re.I):
        p = placa_ok(m.group(1))
        if p:
            scored.append((1, p))
    for m in re.finditer(
        r"carro\s*#?\s*[A-Z0-9\-]+\s*[^\n]{0,30}?([A-Z]{2,3}\d{3,4})\s*\|",
        t,
        re.I,
    ):
        p = placa_ok(m.group(1))
        if p:
            scored.append((0, p))
    for m in re.finditer(r"([A-Z]{2,3}\d{4})\s*\|", t.upper()):
        p = placa_ok(m.group(1))
        if p:
            scored.append((1, p))
    scored.sort()
    return scored[0][1] if scored else None


SKIP_SHEET = re.compile(r"VENDID|PERDIDA|MOTO|TRASPASO", re.I)


def leer_hoja_vida() -> dict[str, dict]:
    if not HOJA.exists():
        raise SystemExit(f"No encuentro el Excel: {HOJA}")
    wb = openpyxl.load_workbook(HOJA, data_only=True, read_only=True)
    fichas: dict[str, dict] = {}

    for name in wb.sheetnames:
        if SKIP_SHEET.search(name):
            continue
        m = re.search(r"CARRO\s+([A-Z]?-?\d+)", name.upper())
        if not m:
            continue
        key = canon_num(m.group(1))
        ws = wb[name]
        title = name.strip()
        marca = modelo = anio = placa = None
        km = None
        # Solo título + primeras filas (ficha de alta), no toda la bitácora
        fuentes_ficha = [name]

        for i, row in enumerate(ws.iter_rows(max_row=12, max_col=4, values_only=True), 1):
            for cell in row:
                if cell is None:
                    continue
                s = str(cell)
                if i == 1 and isinstance(cell, str) and "CARRO" in cell.upper():
                    title = cell.strip()
                    fuentes_ficha.append(title)
                if i <= 4:
                    # ficha / alta / contrato corto
                    if i <= 2 or re.search(
                        r"CARRO|KIA|GRAND|SOLUTO|HYUNDAI|PLACA|CHASIS|👉🏼|👉",
                        s,
                        re.I,
                    ):
                        fuentes_ficha.append(s[:1200])
                if placa is None and i <= 6:
                    p = extract_placa(s)
                    if p:
                        placa = p
            if len(row) > 2 and row[2] is not None:
                try:
                    v = float(row[2])
                    if 50 < v < 800_000:
                        km = int(v)
                except Exception:
                    pass

        # km más reciente (columna C)
        for row in ws.iter_rows(min_row=13, max_col=3, values_only=True):
            if len(row) > 2 and row[2] is not None:
                try:
                    v = float(row[2])
                    if 50 < v < 800_000:
                        km = int(v)
                except Exception:
                    pass

        for src in fuentes_ficha:
            ma, mo, an = parse_ficha(src)
            marca = marca or ma
            modelo = modelo or mo
            anio = anio or an

        # título manda
        ma, mo, an = parse_ficha(title)
        if ma:
            marca = ma
        if mo:
            modelo = mo
        if an:
            anio = an
        ma, mo, an = parse_ficha(name)
        if ma:
            marca = ma
        if mo:
            modelo = mo
        if an:
            anio = an

        # Si hay Kia/Soluto o Grand i10 en título sin año, no inventar
        rec = {
            "numero": key,
            "sheet": name,
            "title": title[:120],
            "marca": marca,
            "modelo": modelo,
            "anio": anio,
            "placa": placa,
            "km": km,
            "emp_hint": "GOLD" if key.startswith("G") else None,
        }
        prev = fichas.get(key)
        if prev:
            for f in ("marca", "modelo", "anio", "placa", "km"):
                if not prev.get(f) and rec.get(f):
                    prev[f] = rec[f]
            if len(rec["title"]) > len(prev.get("title") or ""):
                prev["title"] = rec["title"]
        else:
            fichas[key] = rec

    return fichas


def main() -> None:
    fichas = leer_hoja_vida()
    print(f"Hojas útiles: {len(fichas)}")
    print(
        "  marca",
        sum(1 for v in fichas.values() if v["marca"]),
        "| modelo",
        sum(1 for v in fichas.values() if v["modelo"]),
        "| año",
        sum(1 for v in fichas.values() if v["anio"]),
        "| placa",
        sum(1 for v in fichas.values() if v["placa"]),
    )
    print("  marcas", dict(Counter(v["marca"] for v in fichas.values())))
    print("  modelos", dict(Counter(v["modelo"] for v in fichas.values())))

    veh = sb_get(
        "vehiculos?select=id,numero,placa,marca,modelo,anio,km_actual,estado,"
        "empresa:empresas(codigo)&limit=1000"
    )
    print(f"\nVehículos en DB: {len(veh)}")

    updates = []
    sin_hoja = []
    for v in veh:
        ck = canon_num(v["numero"])
        f = fichas.get(ck) or fichas.get("G" + ck)
        if not f:
            sin_hoja.append(f"{(v.get('empresa') or {}).get('codigo')}·{v['numero']}")
            continue
        patch: dict = {}
        # Completar vacíos; no pisar datos ya cargados
        if f["marca"] and not (v.get("marca") or "").strip():
            patch["marca"] = f["marca"]
        if f["modelo"] and not (v.get("modelo") or "").strip():
            patch["modelo"] = f["modelo"]
        if f["anio"] is not None and v.get("anio") is None:
            patch["anio"] = f["anio"]
        if f["placa"] and not (v.get("placa") or "").strip():
            patch["placa"] = f["placa"]
        if f["km"] and (v.get("km_actual") is None or v.get("km_actual") == 0):
            patch["km_actual"] = f["km"]
        if patch:
            updates.append(
                {
                    "id": v["id"],
                    "numero": v["numero"],
                    "emp": (v.get("empresa") or {}).get("codigo"),
                    "patch": patch,
                    "from": f["title"],
                }
            )

    print(f"Match con hoja: {len(veh) - len(sin_hoja)} | sin hoja: {len(sin_hoja)}")
    if sin_hoja[:15]:
        print("  sin hoja (muestra):", ", ".join(sin_hoja[:15]))
    print(f"Updates: {len(updates)}")
    print(
        "  +marca",
        sum(1 for u in updates if "marca" in u["patch"]),
        "+modelo",
        sum(1 for u in updates if "modelo" in u["patch"]),
        "+año",
        sum(1 for u in updates if "anio" in u["patch"]),
        "+placa",
        sum(1 for u in updates if "placa" in u["patch"]),
        "+km",
        sum(1 for u in updates if "km_actual" in u["patch"]),
    )
    print("\nMuestra:")
    for u in updates[:12]:
        print(f"  {u['emp']} · {u['numero']} ← {u['patch']}  ({u['from'][:60]})")

    out = ROOT / "scripts" / ".cache-hv-enriquecer.json"
    out.write_text(json.dumps({"updates": updates, "sin_hoja": sin_hoja}, ensure_ascii=False, indent=2))
    print(f"\nPlan en {out}")

    if not COMMIT:
        print("\nDRY-RUN. Si cuadra: python3 scripts/enriquecer-desde-hv.py --commit")
        return

    ok = 0
    for u in updates:
        if sb_patch("vehiculos", u["id"], u["patch"]):
            ok += 1
    print(f"\nEscritos: {ok}/{len(updates)}")


if __name__ == "__main__":
    main()
